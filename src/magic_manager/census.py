"""Census reconciliation — diff a physical count against owned inventory.

The user counts ALL their loose cards for a set family into an add-mode
master-list XLSX (sheet ``checklist``; only filled ``qty_normal`` / ``qty_foil``
cells count). Products kept OFF the sheet (built decks, set-aside decks, scene
boxes) are named by MTGJSON fileName. ``target = census + Σ product recipes``;
the plan is the difference between that target and OWNED inventory over the
sheet's printings:

* **adds** — target above owned (written as one ``import-block`` event);
* **removals** — owned above target (written as one ``adhoc`` event), never
  below the copies pledged to built decks.

It is a diff, never an overwrite and never an ingest of the sheet itself.
Driven by ``scripts/census_reconcile.py`` (skill ``census-reconcile``).
"""

from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

from . import db, decks, ingest, inventory, mtgjson, sets, util

Key = tuple[str, str]  # (scryfall_id, finish)

_SHEET = "checklist"
_FINISH_COLS = (("qty_normal", "nonfoil"), ("qty_foil", "foil"))
_REQUIRED = ("set", "collector_number", "name") + tuple(c for c, _ in _FINISH_COLS)
_CHUNK = 500


class CensusError(ValueError):
    """Bad input: unreadable sheet, unresolved rows, unknown product, …"""


class CensusRefused(CensusError):
    """The plan would remove copies pledged to built decks; nothing written."""

    def __init__(self, plan: "CensusPlan"):
        self.plan = plan
        super().__init__(
            f"{len(plan.below_pledged)} removal line(s) would drop below the copies "
            "pledged to built decks — nothing written"
        )


@dataclass
class CensusSheet:
    anchor: str
    universe: set[Key]
    census: Counter
    labels: dict[Key, tuple[str, str, str]]  # key -> (set, cn, name)


@dataclass
class CensusLine:
    """One (printing, finish) row of the plan. ``change`` is the magnitude of the
    adjustment (copies to add on an add line, to remove on a removal line)."""

    scryfall_id: str
    finish: str
    set: str
    cn: str
    name: str
    census: int
    products: int
    owned: int
    pledged: int
    target: int
    change: int


@dataclass
class CensusPlan:
    anchor: str
    products: list[str]
    census_copies: int
    product_copies: int
    target_copies: int
    owned_copies: int
    adds: list[CensusLine] = field(default_factory=list)
    removals: list[CensusLine] = field(default_factory=list)
    below_pledged: list[CensusLine] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "anchor": self.anchor,
            "products": list(self.products),
            "totals": {
                "census_copies": self.census_copies,
                "product_copies": self.product_copies,
                "target_copies": self.target_copies,
                "owned_copies": self.owned_copies,
                "add_rows": len(self.adds),
                "add_copies": sum(ln.change for ln in self.adds),
                "removal_rows": len(self.removals),
                "removal_copies": sum(ln.change for ln in self.removals),
                "below_pledged": len(self.below_pledged),
            },
            "adds": [asdict(ln) for ln in self.adds],
            "removals": [asdict(ln) for ln in self.removals],
            "below_pledged": [asdict(ln) for ln in self.below_pledged],
        }


# ---------- reading the sheet ----------


def _qty(value, where: str) -> int:
    """A filled cell as a non-negative int (floats like 3.0 from Excel are fine)."""
    if isinstance(value, bool):
        raise CensusError(f"{where}: not a non-negative whole number: {value!r}")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, str) and value.strip().isdigit():
        value = int(value.strip())
    if not isinstance(value, int) or value < 0:
        raise CensusError(f"{where}: not a non-negative whole number: {value!r}")
    return value


def read_census(path: str | Path, *, conn=None) -> CensusSheet:
    """Read a filled add-mode master-list XLSX into a :class:`CensusSheet`."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise CensusError("census needs the XLSX checklist")
    if not path.exists():
        raise CensusError(f"file not found: {path}")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
    except Exception as e:  # corrupt / not a workbook
        raise CensusError(f"cannot read {path.name}: {e}") from e
    try:
        if _SHEET not in wb.sheetnames:
            raise CensusError(f"{path.name}: no '{_SHEET}' sheet — is this a master-list checklist?")
        rows = list(wb[_SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise CensusError(f"{path.name}: the '{_SHEET}' sheet is empty")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c in _REQUIRED if c not in header]
    if missing:
        raise CensusError(f"{path.name}: missing column(s) {', '.join(missing)}")

    parsed: list[tuple[int, str, str, str, dict[str, object]]] = []
    for n, raw in enumerate(rows[1:], start=2):
        d = dict(zip(header, raw))
        set_code = d.get("set")
        cn = d.get("collector_number")
        if set_code in (None, "") and cn in (None, ""):
            continue  # blank spacer row
        cn_s = str(int(cn)) if isinstance(cn, float) and cn.is_integer() else str(cn).strip()
        parsed.append((n, str(set_code).strip().lower(), cn_s, str(d.get("name") or ""),
                       {c: d.get(c) for c, _ in _FINISH_COLS}))
    if not parsed:
        raise CensusError(f"{path.name}: no card rows")

    meta = sets.read_master_list_meta(path) or {}
    anchor = str(meta.get("anchor_code") or parsed[0][1]).lower()

    universe: set[Key] = set()
    census: Counter = Counter()
    labels: dict[Key, tuple[str, str, str]] = {}
    unresolved: list[tuple[int, str, str, str]] = []
    with db.transaction(conn) as c:
        for n, set_code, cn, name, cells in parsed:
            row = c.execute(
                "SELECT scryfall_id FROM cards WHERE set_code = ? AND collector_number = ?",
                (set_code, cn),
            ).fetchone()
            if row is None:
                unresolved.append((n, set_code, cn, name))
                continue
            sid = row["scryfall_id"]
            for col, finish in _FINISH_COLS:
                key = (sid, finish)
                universe.add(key)
                labels[key] = (set_code, cn, name)
                cell = cells[col]
                if cell not in (None, ""):
                    census[key] += _qty(cell, f"row {n} {set_code} {cn} {name} [{col}]")
    if unresolved:
        shown = [f"row {n}: {s} {cn} {name}".rstrip() for n, s, cn, name in unresolved[:25]]
        more = f"\n  …and {len(unresolved) - 25} more" if len(unresolved) > 25 else ""
        codes = sorted({s for _, s, _, _ in unresolved})
        raise CensusError(
            f"{len(unresolved)} row(s) not in the local catalog:\n  " + "\n  ".join(shown) + more
            + "\nsync the set(s) first: " + "; ".join(f"uv run mm set sync {c}" for c in codes)
        )
    return CensusSheet(anchor=anchor, universe=universe, census=census, labels=labels)


# ---------- planning ----------


def _product_needs(products: list[str]) -> Counter:
    off: Counter = Counter()
    for file_name in products:
        try:
            needs = decks.precon_recipe_needs(file_name)
        except mtgjson.MtgJsonError as e:
            raise CensusError(f"unknown or unreadable product {file_name!r}: {e}") from e
        if not needs:
            raise CensusError(f"product {file_name!r} has no recipe cards")
        off.update(needs)
    return off


def _read_counts(conn, sids: list[str]) -> tuple[Counter, Counter]:
    owned: Counter = Counter()
    pledged: Counter = Counter()
    for i in range(0, len(sids), _CHUNK):
        part = sids[i:i + _CHUNK]
        marks = ",".join("?" * len(part))
        for r in conn.execute(
            f"SELECT scryfall_id, finish, quantity FROM inventory WHERE scryfall_id IN ({marks})",
            part,
        ):
            owned[(r["scryfall_id"], r["finish"])] = r["quantity"]
        for r in conn.execute(
            f"SELECT scryfall_id, finish, SUM(count) AS n FROM deck_assignments "
            f"WHERE scryfall_id IN ({marks}) GROUP BY scryfall_id, finish",
            part,
        ):
            pledged[(r["scryfall_id"], r["finish"])] = r["n"]
    return owned, pledged


def _build(sheet: CensusSheet, products: list[str], off: Counter, conn) -> CensusPlan:
    outside = sorted(k for k in off if k not in sheet.universe)
    if outside:
        raise CensusError(
            f"{len(outside)} product card(s) are not on the sheet (the products must "
            "be the same set family as the sheet): "
            + ", ".join(f"{sid} {fin}" for sid, fin in outside[:10])
            + (" …" if len(outside) > 10 else "")
        )
    owned, pledged = _read_counts(conn, sorted({sid for sid, _ in sheet.universe}))

    adds: list[CensusLine] = []
    removals: list[CensusLine] = []
    target_total = 0
    for key in sheet.universe:
        tgt = sheet.census[key] + off[key]
        inv = owned[key]
        target_total += tgt
        if tgt == inv:
            continue
        set_code, cn, name = sheet.labels[key]
        line = CensusLine(
            scryfall_id=key[0], finish=key[1], set=set_code, cn=cn, name=name,
            census=sheet.census[key], products=off[key], owned=inv,
            pledged=pledged[key], target=tgt, change=abs(tgt - inv),
        )
        (adds if tgt > inv else removals).append(line)

    def order(ln: CensusLine):
        return (ln.set, util.cn_sort_key(ln.cn), ln.finish)

    adds.sort(key=order)
    removals.sort(key=order)
    return CensusPlan(
        anchor=sheet.anchor,
        products=list(products),
        census_copies=sum(sheet.census.values()),
        product_copies=sum(off.values()),
        target_copies=target_total,
        owned_copies=sum(owned[k] for k in sheet.universe),
        adds=adds,
        removals=removals,
        below_pledged=[ln for ln in removals if ln.owned - ln.change < ln.pledged],
    )


def plan(path: str | Path, products: list[str] | None = None, *, conn=None) -> CensusPlan:
    """Diff ``census + Σ products`` against owned inventory (read-only).

    ``products`` are MTGJSON fileNames and may repeat (two copies = listed twice).
    """
    products = list(products or [])
    off = _product_needs(products)
    with db.transaction(conn) as c:
        sheet = read_census(path, conn=c)
        return _build(sheet, products, off, c)


# ---------- applying ----------


def apply(path: str | Path, products: list[str] | None = None) -> dict:
    """Write the plan in ONE transaction: adds as an ``import-block`` event,
    removals as an ``adhoc`` event. Refuses (nothing written) if any removal would
    drop below the copies pledged to built decks."""
    products = list(products or [])
    off = _product_needs(products)  # network/disk reads stay outside the write transaction
    with db.transaction() as c:
        sheet = read_census(path, conn=c)
        p = _build(sheet, products, off, c)
        if p.below_pledged:
            raise CensusRefused(p)

        adds_event = removals_event = None
        touched: set[Key] = set()
        if p.adds:
            label = f"census:{p.anchor} loose cards"
            with ingest.open_ingest_event("import-block", label=label, conn=c) as rec:
                adds_event = rec.ingest_id
                for ln in p.adds:
                    inventory.inventory_add(ln.scryfall_id, ln.finish, ln.change,
                                            conn=c, ingest_id=rec.ingest_id)
                    # inventory_add already appended the ledger delta; only tally.
                    if ln.owned == 0:
                        rec.rows_added += 1
                    else:
                        rec.rows_updated += 1
                    touched.add((ln.scryfall_id, ln.finish))
        if p.removals:
            label = f"census:{p.anchor} over-count removal"
            with ingest.open_ingest_event("adhoc", label=label, conn=c) as rec:
                removals_event = rec.ingest_id
                for ln in p.removals:
                    inventory.inventory_remove(ln.scryfall_id, ln.finish, ln.change,
                                               conn=c, ingest_id=rec.ingest_id)
                    if ln.owned - ln.change == 0:
                        rec.rows_zeroed += 1
                    else:
                        rec.rows_updated += 1
                    touched.add((ln.scryfall_id, ln.finish))

        after = _build(sheet, products, off, c)
        if after.adds or after.removals:
            raise RuntimeError(
                f"census apply did not converge: {len(after.adds)} add(s) and "
                f"{len(after.removals)} removal(s) remain"
            )
        drift = [d for d in ingest.reconcile_inventory_ledger(c)
                 if (d["scryfall_id"], d["finish"]) in touched]
        if drift:
            raise RuntimeError(f"inventory and ledger disagree after census apply: {drift[:5]}")
        return {
            "adds_event": adds_event,
            "removals_event": removals_event,
            "added_copies": sum(ln.change for ln in p.adds),
            "removed_copies": sum(ln.change for ln in p.removals),
        }
