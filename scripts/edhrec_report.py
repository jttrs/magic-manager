"""Deterministic EDHREC report generator.

Three workflows, one script (subcommand-dispatched, the single source of truth
for all EDHREC report math + formatting — the CLI/skills are thin relays):

  commander <card>   Workflow A — cards with the highest inclusion rate when
                     <card> is the commander.
  card <card>        Workflow B — the most common commanders that run <card> in
                     the 99 (not as the commander).
  rankings <scope>   Workflow C — general rankings: 'commanders' (by deck count),
                     'cards' (top played), or 'salt' (saltiest).

<card> is a user-facing reference — a NAME ("Sol Ring") or a printing
("SET CN", e.g. "cmm 425"). It's up-leveled to its oracle name via Scryfall
(all reprints collapse to one EDHREC page) before the EDHREC slug is derived.

Every report emits (matching sealed-value / construct-value):
  * a markdown table to stdout (Scryfall-hyperlinked names) for chat relay, and
  * JSON + XLSX artifacts under output/edhrec/reports/.

Usage:
    uv run python scripts/edhrec_report.py commander "Atraxa, Praetors' Voice" --top 25
    uv run python scripts/edhrec_report.py card "Sol Ring"
    uv run python scripts/edhrec_report.py rankings commanders --timeframe week
    uv run python scripts/edhrec_report.py rankings salt

Exit codes:
    0 — report written
    2 — bad invocation / card not found / EDHREC page missing
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
_OUTPUT_TYPE = "edhrec"  # → output/edhrec/reports/
sys.path.insert(0, str(ROOT / "src"))

from magic_manager import edhrec, sets, util  # noqa: E402


# ---------- rendering ----------

def _scryfall_search_url(name: str) -> str:
    from urllib.parse import quote
    return f"https://scryfall.com/search?q={quote('!\"' + name + '\"')}"


def _md_commander(res, top: int, prices_note: str) -> str:
    rows = [r for r in res.rows if r.list_tag == "topcards"]
    rows.sort(key=lambda r: -(r.inclusion_pct or 0))
    rows = rows[:top]
    out = [
        f"## EDHREC — top cards in **{res.name}** decks",
        "",
        f"Highest inclusion rate across {res.name} decks (workflow A). {prices_note}",
        "",
        "| # | Card | Incl % | Decks | Synergy | Type | MV | Lowest $ |",
        "|--:|------|-------:|------:|--------:|------|---:|---------:|",
    ]
    for i, r in enumerate(rows, 1):
        out.append(
            f"| {i} | [{r.name}]({_scryfall_search_url(r.name)}) "
            f"| {r.inclusion_pct if r.inclusion_pct is not None else '—'} "
            f"| {r.num_decks or '—'} "
            f"| {round(r.synergy, 3) if r.synergy is not None else '—'} "
            f"| {r.type_line or '—'} "
            f"| {int(r.cmc) if r.cmc is not None else '—'} "
            f"| {util.fmt_usd(r.lowest_usd)} |"
        )
    return "\n".join(out)


def _md_card(res, top: int, prices_note: str) -> str:
    rows = [r for r in res.rows if r.list_tag == "topcommanders"]
    rows.sort(key=lambda r: -(r.num_decks or 0))
    rows = rows[:top]
    out = [
        f"## EDHREC — top commanders running **{res.name}**",
        "",
        f"Most common commanders whose decks include {res.name} in the 99 (workflow B). {prices_note}",
        "",
        "| # | Commander | Decks | Incl % | Type | MV | Lowest $ |",
        "|--:|-----------|------:|-------:|------|---:|---------:|",
    ]
    for i, r in enumerate(rows, 1):
        out.append(
            f"| {i} | [{r.name}]({_scryfall_search_url(r.name)}) "
            f"| {r.num_decks or '—'} "
            f"| {r.inclusion_pct if r.inclusion_pct is not None else '—'} "
            f"| {r.type_line or '—'} "
            f"| {int(r.cmc) if r.cmc is not None else '—'} "
            f"| {util.fmt_usd(r.lowest_usd)} |"
        )
    return "\n".join(out)


def _md_rankings(res, top: int, prices_note: str) -> str:
    rows = res.rows[:top]
    is_salt = res.scope == "salt"
    metric_hdr = "Salt" if is_salt else "Decks"
    entity_hdr = "Commander" if res.scope == "commanders" else "Card"
    out = [
        f"## EDHREC rankings — {res.scope} ({res.timeframe})",
        "",
        f"General {res.scope} ranking (workflow C). {prices_note}",
        "",
        f"| Rank | {entity_hdr} | {metric_hdr} | Type | MV | Lowest $ |",
        "|-----:|------------|------:|------|---:|---------:|",
    ]
    for r in rows:
        metric = r.salt if is_salt else r.num_decks
        metric_s = (round(metric, 2) if is_salt and metric is not None
                    else (metric if metric is not None else "—"))
        out.append(
            f"| {r.rank} | [{r.name}]({_scryfall_search_url(r.name)}) "
            f"| {metric_s} "
            f"| {r.type_line or '—'} "
            f"| {int(r.cmc) if r.cmc is not None else '—'} "
            f"| {util.fmt_usd(r.lowest_usd)} |"
        )
    return "\n".join(out)


# ---------- artifacts ----------

def _row_dict(r) -> dict:
    return {
        "name": r.name, "slug": r.slug, "oracle_id": r.oracle_id,
        "list_tag": r.list_tag, "num_decks": r.num_decks,
        "potential_decks": r.potential_decks, "inclusion_pct": r.inclusion_pct,
        "synergy": r.synergy, "lift": r.lift, "trend_zscore": r.trend_zscore,
        "salt": r.salt, "rank": r.rank, "type_line": r.type_line, "cmc": r.cmc,
        "mana_cost": r.mana_cost, "color_identity": r.color_identity,
        "rarity": r.rarity, "lowest_usd": r.lowest_usd,
        "lowest_usd_foil": r.lowest_usd_foil,
    }


def _write_json(res, kind: str, out_path: Path, prices_as_of: str | None) -> None:
    payload = {
        "kind": kind,
        "slug": res.slug,
        "name": res.name,
        "scope": res.scope,
        "timeframe": res.timeframe,
        "prices_as_of": prices_as_of,
        "generated_at": datetime.now(UTC).isoformat(),
        "rows": [_row_dict(r) for r in res.rows],
    }
    out_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _write_xlsx(res, kind: str, out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = kind
    headers = ["list_tag", "name", "num_decks", "potential_decks", "inclusion_pct",
               "synergy", "lift", "salt", "rank", "type_line", "mana_value",
               "lowest_usd", "lowest_usd_foil", "oracle_id"]
    ws.append(headers)
    for r in res.rows:
        ws.append([
            r.list_tag, r.name, r.num_decks, r.potential_decks, r.inclusion_pct,
            r.synergy, r.lift, r.salt, r.rank, r.type_line,
            int(r.cmc) if r.cmc is not None else None,
            r.lowest_usd, r.lowest_usd_foil, r.oracle_id,
        ])
    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row=row_idx, column=12).number_format = '"$"#,##0.00'
        ws.cell(row=row_idx, column=13).number_format = '"$"#,##0.00'
    widths = {1: 16, 2: 34, 3: 10, 4: 12, 5: 11, 6: 9, 7: 8, 8: 8, 9: 6,
              10: 30, 11: 6, 12: 11, 13: 13, 14: 38}
    for ci, w in widths.items():
        ws.column_dimensions[get_column_letter(ci)].width = w
    for col in range(1, ws.max_column + 1):
        ws.cell(row=1, column=col).font = Font(bold=True)
        ws.cell(row=1, column=col).alignment = Alignment(horizontal="left")
    ws.freeze_panes = "A2"
    util.apply_base_font_size(ws)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ---------- driver ----------

def _prices_note(oids: list[str], *, refresh: bool) -> str:
    """Ensure prices for the resolved oracle set, return a 'prices as of' note."""
    # Resolve the sets those oracle cards belong to, ensure fresh, then report.
    if not oids:
        return ""
    with sets.db.connect() as c:
        placeholders = ",".join("?" for _ in oids)
        rows = c.execute(
            f"SELECT DISTINCT set_code, scryfall_id FROM cards WHERE oracle_id IN ({placeholders})",
            oids,
        ).fetchall()
    codes = sorted({r[0] for r in rows})
    sids = [r[1] for r in rows]
    if codes:
        sets.ensure_priced(codes, refresh_stale=refresh, log=lambda m: print(m, file=sys.stderr))
    newest, oldest = sets.prices_as_of(sids)
    if newest:
        return f"Prices fetched: {newest}." if newest == oldest else f"Prices fetched: {oldest}–{newest}."
    return "Prices: local (best-effort)."


def main() -> int:
    # Shared options live on a parent parser so they can appear AFTER the
    # subcommand (argparse won't accept a top-level flag post-subcommand).
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--no-refresh", action="store_true",
                        help="Use local prices as-is; don't re-sync stale sets.")

    ap = argparse.ArgumentParser(description="EDHREC report generator.")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_cmd = sub.add_parser("commander", parents=[common],
                           help="Workflow A: top cards for a commander.")
    p_cmd.add_argument("card")
    p_cmd.add_argument("--top", type=int, default=25)

    p_card = sub.add_parser("card", parents=[common],
                            help="Workflow B: top commanders running a card.")
    p_card.add_argument("card")
    p_card.add_argument("--top", type=int, default=25)

    p_rank = sub.add_parser("rankings", parents=[common],
                            help="Workflow C: general rankings.")
    p_rank.add_argument("scope", choices=["commanders", "cards", "salt"])
    p_rank.add_argument("--timeframe", default="week")
    p_rank.add_argument("--top", type=int, default=50)

    args = ap.parse_args()
    refresh = not args.no_refresh
    ts = datetime.now(UTC).strftime("%Y-%m-%d-%H%M%S")
    out_dir = util.output_dir(_OUTPUT_TYPE, "reports")

    try:
        if args.cmd == "commander":
            res = edhrec.sync_commander(edhrec.resolve_oracle_name(args.card))
            renderer, base = _md_commander, f"{res.slug}-commander-{ts}"
        elif args.cmd == "card":
            res = edhrec.sync_card(edhrec.resolve_oracle_name(args.card))
            renderer, base = _md_card, f"{res.slug}-card-{ts}"
        else:  # rankings
            res = edhrec.sync_rankings(args.scope, args.timeframe)
            renderer, base = _md_rankings, f"rankings-{res.scope}-{res.timeframe}-{ts}"
    except edhrec.EdhrecError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2

    # sync -> ensure prices for the resolved oracle set -> enrich ONCE (so the
    # local metadata reflects any freshly-synced prices) -> render.
    oids = [r.oracle_id for r in res.rows if r.oracle_id]
    note = _prices_note(oids, refresh=refresh)
    edhrec.enrich_rows(res.rows)
    md = renderer(res, args.top, note)

    json_path = out_dir / f"{base}.json"
    xlsx_path = out_dir / f"{base}.xlsx"
    _write_json(res, args.cmd, json_path, note or None)
    _write_xlsx(res, args.cmd, xlsx_path)

    print(md)
    print()
    print(f"→ {json_path}")
    print(f"→ {xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
