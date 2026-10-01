"""Golden-ish tests for magic_manager.exports.xlsx — the shared workbook writer.

Pins the conventions every migrated caller relies on: bold+left headers, money
("$"#,##0.00) and text (@) number formats, per-column widths, per-column
hyperlinks (via the (row, index) callable), hidden _meta sheet, freeze panes,
and the 16pt base font applied to every sheet. Reads the written file back with
openpyxl so a future change to the writer that drifts any of these fails here.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from magic_manager import util
from magic_manager.exports import xlsx


def _write(tmp_path, specs) -> Path:
    out = tmp_path / "wb.xlsx"
    xlsx.write_workbook(out, specs)
    return out


def test_single_sheet_formats_widths_and_font(tmp_path):
    spec = xlsx.SheetSpec(
        title="results",
        headers=["set", "cn", "name", "usd"],
        rows=[["tla", "5", "Card A", 1.5], ["tla", "6", "Card B", 2.0]],
        widths={1: 6, 3: 40},
        money_cols=(4,),
        text_cols=(2,),
    )
    wb = openpyxl.load_workbook(_write(tmp_path, [spec]))
    assert wb.sheetnames == ["results"]
    ws = wb["results"]
    assert [c.value for c in ws[1]] == ["set", "cn", "name", "usd"]
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.cell(row=1, column=1).alignment.horizontal == "left"
    assert ws.cell(row=2, column=4).number_format == '"$"#,##0.00'
    assert ws.cell(row=2, column=2).number_format == "@"
    assert ws.column_dimensions["A"].width == 6
    assert ws.column_dimensions["C"].width == 40
    assert ws.freeze_panes == "A2"
    # base font bumped to 16pt on populated cells
    assert ws.cell(row=2, column=1).font.size == util.XLSX_FONT_SIZE


def test_hyperlinks_use_row_and_index(tmp_path):
    uris = ["https://scryfall.com/a", None]  # second row intentionally has no link
    spec = xlsx.SheetSpec(
        title="r", headers=["name", "x"],
        rows=[["A", 1], ["B", 2]],
        hyperlinks={1: lambda _row, i: uris[i]},
    )
    wb = openpyxl.load_workbook(_write(tmp_path, [spec]))
    ws = wb["r"]
    assert ws.cell(row=2, column=1).hyperlink.target == "https://scryfall.com/a"
    assert ws.cell(row=3, column=1).hyperlink is None  # None url → no link
    # linked cell is styled blue + underline
    assert ws.cell(row=2, column=1).font.underline == "single"


def test_hyperlink_callable_can_read_row_values(tmp_path):
    spec = xlsx.SheetSpec(
        title="r", headers=["set", "cn"],
        rows=[["tla", "5"]],
        hyperlinks={2: lambda row, _i: f"https://x/{row[0]}/{row[1]}"},
    )
    wb = openpyxl.load_workbook(_write(tmp_path, [spec]))
    assert wb["r"].cell(row=2, column=2).hyperlink.target == "https://x/tla/5"


def test_multi_sheet_order_and_meta_hidden(tmp_path):
    a = xlsx.SheetSpec(title="tree", headers=["x"], rows=[[1]])
    b = xlsx.SheetSpec(title="sheets", headers=["y"], rows=[[2]])
    meta = xlsx.meta_sheet([("kind", "sealed"), ("n", "3")])
    wb = openpyxl.load_workbook(_write(tmp_path, [a, b, meta]))
    assert wb.sheetnames == ["tree", "sheets", "_meta"]
    assert wb["_meta"].sheet_state == "hidden"
    assert [(r[0].value, r[1].value) for r in wb["_meta"].iter_rows()] == \
        [("key", "value"), ("kind", "sealed"), ("n", "3")]
    # a non-frozen meta sheet
    assert wb["_meta"].freeze_panes is None


def test_freeze_can_be_disabled(tmp_path):
    spec = xlsx.SheetSpec(title="r", headers=["x"], rows=[[1]], freeze=None)
    wb = openpyxl.load_workbook(_write(tmp_path, [spec]))
    assert wb["r"].freeze_panes is None


def test_mkdir_creates_missing_parents(tmp_path):
    nested = tmp_path / "deep" / "dir" / "out.xlsx"
    xlsx.write_workbook(nested, [xlsx.SheetSpec(title="r", headers=["x"], rows=[[1]])])
    assert nested.exists()
