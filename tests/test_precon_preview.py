"""Regression tests for the precon deck-checklist PREVIEW
(`sets._summarize_deck_checklist`) — it must report the SAME absolute-count
vocabulary the ingest (`_apply_precon_checklist`) uses, so a modify row that
keeps an existing built copy and adds a deconstructed one can't be misread as
"constructed=0".

The 2026-09-15 bug: the preview stuffed the signed DELTA into fields named
`constructed_qty`/`deconstructed_qty`, so a FIC commander at built=1 edited to
(constructed=1, deconstructed=1) previewed as `constructed_qty:0` (delta) while
ingest correctly applied count_after=(1,1). This pins preview == ingest.

Offline: tmp_db + fake_mtgjson (import_precon's deck fetch) + a hand-built XLSX.
"""

from __future__ import annotations

from openpyxl import Workbook


def _write_modify_checklist(path, rows):
    """Minimal precon MODIFY checklist XLSX. `rows`: list of
    (file_name, deck_name, constructed_qty, deconstructed_qty)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "checklist"
    ws.append(["file_name", "set", "deck_name", "color", "type", "release_date",
               "commander", "card_count", "usd_total",
               "constructed_qty", "deconstructed_qty"])
    for fn, name, c, d in rows:
        ws.append([fn, "tmc", name, "WUB", "Commander Deck", "2025-01-01",
                   "", 100, 10.0, c, d])
    meta = wb.create_sheet("_meta")
    meta.append(["key", "value"])
    for k, v in (("kind", "precon"), ("mode", "modify"), ("slug", "precons")):
        meta.append([k, v])
    wb.save(str(path))


def _seed_built_precon(fake_mtgjson, make_card, make_precon_deck, fake_scryfall,
                       file_name, name):
    """Create a built precon deck (built=1, decon=0) via import_precon."""
    from magic_manager import decks
    sid = f"{file_name}-c1"
    fake_mtgjson(deck=make_precon_deck(
        name, "Commander Deck",
        [{"sid": sid, "name": "Cmdr", "set": "tmc", "cn": "1", "count": 1, "board": "commander"}]))
    fake_scryfall(search=[make_card(id=sid, set="tmc", collector_number="1", name="Cmdr")])
    decks.import_precon(file_name)


def test_modify_preview_matches_ingest_keep_plus_add(
        tmp_db, tmp_path, fake_scryfall, fake_mtgjson, make_card, make_precon_deck):
    """The exact FIC case: built=1 already owned, edit to (constructed=1,
    deconstructed=1). Preview must show count_before=[1,0], count_after=[1,1],
    constructed_qty=1 (NOT 0), delta=[0,1] — and equal the ingest's per_row."""
    from magic_manager import sets as sets_mod, parsers
    _seed_built_precon(fake_mtgjson, make_card, make_precon_deck, fake_scryfall,
                       "ScionsSpellcraft_FIC", "Scions & Spellcraft")

    path = tmp_path / "precons-modify-checklist.xlsx"
    _write_modify_checklist(path, [("ScionsSpellcraft_FIC", "Scions & Spellcraft", 1, 1)])

    # PREVIEW
    summary = sets_mod._summarize_deck_checklist(path, {"kind": "precon", "mode": "modify"})
    assert summary["kind"] == "precon" and summary["mode"] == "modify"
    assert len(summary["filled"]) == 1
    row = summary["filled"][0]
    assert row["count_before"] == [1, 0]
    assert row["count_after"] == [1, 1]
    assert row["constructed_qty"] == 1, "resulting built count must be 1, not the delta 0"
    assert row["deconstructed_qty"] == 1
    assert tuple(row["delta"]) == (0, 1)

    # INGEST — same file, must agree on before/after/delta
    parsed = parsers.parse_jumpstart_list_xlsx(path)
    ing = sets_mod._apply_precon_checklist(parsed, mode="modify")
    per = ing["per_row"][0]
    assert tuple(per["count_before"]) == tuple(row["count_before"])
    assert tuple(per["count_after"]) == tuple(row["count_after"])
    assert tuple(per["delta"]) == tuple(row["delta"])


def _write_add_checklist(path, rows):
    """Minimal precon ADD checklist XLSX. `rows`: (file_name, deck_name, acquired_qty)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "checklist"
    ws.append(["file_name", "set", "deck_name", "color", "type", "release_date",
               "commander", "card_count", "usd_total", "acquired_qty"])
    for fn, name, a in rows:
        ws.append([fn, "tmc", name, "WUB", "Commander Deck", "2025-01-01", "", 100, 10.0, a])
    meta = wb.create_sheet("_meta")
    meta.append(["key", "value"])
    for k, v in (("kind", "precon"), ("mode", "add"), ("slug", "precons")):
        meta.append([k, v])
    wb.save(str(path))


def test_add_preview_splits_into_resulting_counts(
        tmp_db, tmp_path, fake_scryfall, fake_mtgjson, make_card, make_precon_deck):
    """ADD mode: acquiring 2 of a net-new buildable precon splits to 1 built +
    1 deconstructed → count_after=[1,1], delta=[1,1] (resulting counts, matching
    ingest's split), not raw acquired_qty in the per-state fields."""
    from magic_manager import sets as sets_mod, parsers
    # net-new: nothing owned yet. Prime mtgjson so default_precon_state='built'.
    fake_mtgjson(deck=make_precon_deck(
        "Fresh Deck", "Commander Deck",
        [{"sid": "fresh-c1", "name": "Cmdr", "set": "tmc", "cn": "1", "count": 1, "board": "commander"}]))
    fake_scryfall(search=[make_card(id="fresh-c1", set="tmc", collector_number="1", name="Cmdr")])

    path = tmp_path / "precons-add-checklist.xlsx"
    _write_add_checklist(path, [("Fresh_TMC", "Fresh Deck", 2)])
    summary = sets_mod._summarize_deck_checklist(path, {"kind": "precon", "mode": "add"})
    assert len(summary["filled"]) == 1
    row = summary["filled"][0]
    # acquired 2 of a net-new buildable precon → split 1 built + 1 deconstructed.
    # The per-state fields are RESULTING counts (count_after), not raw acquired_qty.
    assert row["acquired_qty"] == 2
    assert row["count_before"] == [0, 0]
    assert row["count_after"] == [1, 1]
    assert row["constructed_qty"] == 1 and row["deconstructed_qty"] == 1
    assert tuple(row["delta"]) == (1, 1)
    # (add-mode INGEST goes through _apply_acquired_checklist, which reports a
    # different built/torn_down shape — parity with _apply_precon_checklist only
    # applies to modify mode, covered above.)


def test_modify_preview_untouched_row_skipped(
        tmp_db, tmp_path, fake_scryfall, fake_mtgjson, make_card, make_precon_deck):
    """A row whose entered counts equal current (built=1 → constructed=1,
    deconstructed=0) is a no-op → not in filled[]."""
    from magic_manager import sets as sets_mod
    _seed_built_precon(fake_mtgjson, make_card, make_precon_deck, fake_scryfall,
                       "NoOp_FIC", "No Op")
    path = tmp_path / "precons-modify-checklist.xlsx"
    _write_modify_checklist(path, [("NoOp_FIC", "No Op", 1, 0)])
    summary = sets_mod._summarize_deck_checklist(path, {"kind": "precon", "mode": "modify"})
    assert summary["filled"] == []


def test_modify_preview_lowered_count_surfaces(
        tmp_db, tmp_path, fake_scryfall, fake_mtgjson, make_card, make_precon_deck):
    """Lowering below current (built=1 → constructed=0) shows a negative delta and
    count_after < count_before, so the preview surfaces the will-warn case rather
    than hiding it (ingest won't actually lower — that's an explicit deck delete)."""
    from magic_manager import sets as sets_mod
    _seed_built_precon(fake_mtgjson, make_card, make_precon_deck, fake_scryfall,
                       "Lowered_FIC", "Lowered")
    path = tmp_path / "precons-modify-checklist.xlsx"
    _write_modify_checklist(path, [("Lowered_FIC", "Lowered", 0, 0)])
    summary = sets_mod._summarize_deck_checklist(path, {"kind": "precon", "mode": "modify"})
    assert len(summary["filled"]) == 1
    row = summary["filled"][0]
    assert row["count_before"] == [1, 0]
    assert row["count_after"] == [0, 0]
    assert tuple(row["delta"]) == (-1, 0)
